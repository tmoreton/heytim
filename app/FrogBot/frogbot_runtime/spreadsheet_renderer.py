from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Any

from .artifact_content import document_title, validate_content

MAX_ROWS = 5_000
MAX_COLUMNS = 100


def spreadsheet_value(value: str) -> Any:
    stripped = value.strip()
    if not stripped:
        return ""
    if stripped[0] in "=+-@":
        return value
    if re.fullmatch(r"-?(0|[1-9]\d*)", stripped) and not re.fullmatch(
        r"-?0\d+", stripped
    ):
        try:
            return int(stripped)
        except ValueError:
            return value
    if re.fullmatch(r"-?(?:\d+\.\d+|\d+\.\d*[eE][+-]?\d+)", stripped):
        try:
            return float(stripped)
        except ValueError:
            return value
    try:
        return (
            datetime.fromisoformat(stripped)
            if "T" in stripped
            else date.fromisoformat(stripped)
        )
    except ValueError:
        return value


def csv_rows(content: str) -> list[list[str]]:
    normalized = validate_content(content)
    try:
        dialect = csv.Sniffer().sniff(normalized[:8_192], delimiters=",\t;|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(normalized), dialect=dialect))
    if not rows or not any(any(cell.strip() for cell in row) for row in rows):
        raise ValueError("spreadsheet content must include at least one populated cell")
    if len(rows) > MAX_ROWS:
        raise ValueError(f"spreadsheet content may contain at most {MAX_ROWS} rows")
    if max(map(len, rows), default=0) > MAX_COLUMNS:
        raise ValueError(
            f"spreadsheet content may contain at most {MAX_COLUMNS} columns"
        )
    return rows


def render_xlsx(filename: str, content: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rows = csv_rows(content)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = document_title(filename)[:31]
    cell_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    for row_index, row in enumerate(rows, start=1):
        for column_index, raw_value in enumerate(row, start=1):
            value = spreadsheet_value(raw_value)
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(value, date):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = "#,##0.00" if isinstance(value, float) else "#,##0"
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = cell_border

    header_fill = PatternFill("solid", fgColor="007A3D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2" if len(rows) > 1 else None
    if rows and rows[0]:
        sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 24
    for column_index in range(1, max(map(len, rows), default=0) + 1):
        values = [
            str(row[column_index - 1])
            for row in rows
            if len(row) >= column_index
        ]
        width = min(
            45, max(12, max((len(value) for value in values), default=0) + 4)
        )
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    workbook.properties.title = document_title(filename)
    workbook.properties.creator = "FroggyBot"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
