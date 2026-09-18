from __future__ import annotations

import hashlib
import io
import uuid
import zipfile
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from frogbot_runtime import artifacts


class FakeS3:
    def __init__(self) -> None:
        self.requests: list[dict] = []

    def put_object(self, **request) -> None:
        self.requests.append(request)


class ReadableFakeS3(FakeS3):
    def list_objects_v2(self, **request) -> dict:
        matches = [
            {"Key": item["Key"]}
            for item in self.requests
            if item["Key"].startswith(request["Prefix"])
        ]
        return {"Contents": matches[: request["MaxKeys"]]}

    def get_object(self, **request) -> dict:
        stored = next(item for item in self.requests if item["Key"] == request["Key"])
        return {
            "Body": io.BytesIO(stored["Body"]),
            "ContentType": stored["ContentType"],
            "Metadata": stored.get("Metadata", {}),
        }


def _points_png(size: tuple[int, int] = (320, 180)) -> bytes:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, size[0] - 10, size[1] - 10), outline="black", width=2)
    draw.text((25, 50), "12,000 points", fill="black")
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def test_artifact_tool_writes_only_to_the_scoped_user_prefix(monkeypatch) -> None:
    bucket = "frogbot-user-files-123-us-east-1"
    prefix = f"users/{'a' * 64}/artifacts/12345678-1234-1234-1234-123456789012"
    target = FakeS3()
    monkeypatch.setattr(artifacts, "FILES_BUCKET_NAME", bucket)

    save = artifacts.artifact_tool(prefix, client=target)
    result = save("../Quarterly 📊.md", "# Results\n\nEverything is on track.")

    assert result == "Saved Quarterly.md for the user to download."
    request = target.requests[0]
    assert request["Bucket"] == bucket
    assert request["Key"].startswith(f"{prefix}/")
    assert request["Key"].endswith("Quarterly.md")
    assert "ServerSideEncryption" not in request


def test_png_artifact_round_trips_only_inside_its_scoped_prefix(monkeypatch) -> None:
    prefix = f"users/{'a' * 64}/artifacts/12345678-1234-1234-1234-123456789012"
    target = ReadableFakeS3()
    monkeypatch.setattr(artifacts, "FILES_BUCKET_NAME", "files")
    png = _points_png()

    stored = artifacts.put_png_artifact(
        prefix,
        "../Award deal ✈.png",
        png,
        client=target,
        metadata={"source-url": "https://www.delta.com/deals"},
    )
    loaded = artifacts.load_png_artifact(prefix, stored["artifactId"], client=target)

    assert uuid.UUID(stored["artifactId"])
    assert stored["filename"] == "Award deal.png"
    assert loaded["body"] == png
    assert loaded["filename"] == "Award deal.png"
    assert loaded["metadata"]["source-url"] == "https://www.delta.com/deals"
    assert "ServerSideEncryption" not in target.requests[0]
    with pytest.raises(ValueError, match="unavailable"):
        artifacts.load_png_artifact(
            f"users/{'b' * 64}/artifacts/12345678-1234-1234-1234-123456789012",
            stored["artifactId"],
            client=target,
        )


def test_png_artifact_rejects_non_png_content(monkeypatch) -> None:
    prefix = f"users/{'a' * 64}/artifacts/12345678-1234-1234-1234-123456789012"
    monkeypatch.setattr(artifacts, "FILES_BUCKET_NAME", "files")

    with pytest.raises(ValueError, match="PNG"):
        artifacts.put_png_artifact(
            prefix, "award.png", b"not an image", client=FakeS3()
        )


def test_points_screenshot_crops_viewport_and_rejects_blank_png() -> None:
    cropped = artifacts.crop_points_screenshot(
        _points_png((500, 300)), {"x": 5, "y": 5, "width": 330, "height": 190}
    )

    with Image.open(io.BytesIO(cropped)) as image:
        assert image.size == (330, 190)

    blank = Image.new("RGB", (320, 180), "white")
    output = io.BytesIO()
    blank.save(output, format="PNG")
    with pytest.raises(ValueError, match="blank or uniform"):
        artifacts.crop_points_screenshot(
            output.getvalue(), {"x": 0, "y": 0, "width": 320, "height": 180}
        )


def test_artifact_context_rejects_an_unscoped_prefix() -> None:
    with pytest.raises(ValueError, match="prefix"):
        artifacts.artifact_prefix_from_payload(
            {"artifacts": {"prefix": "users/another-user/artifacts/turn"}},
            "a" * 64,
        )


def test_artifact_context_is_bound_to_the_invoking_user() -> None:
    actor_id = "a" * 64
    prefix = f"users/{actor_id}/artifacts/12345678-1234-1234-1234-123456789012"
    assert (
        artifacts.artifact_prefix_from_payload(
            {"artifacts": {"prefix": prefix}}, actor_id
        )
        == prefix
    )
    with pytest.raises(ValueError, match="invoking user"):
        artifacts.artifact_prefix_from_payload(
            {"artifacts": {"prefix": prefix}}, "b" * 64
        )


def test_bot_artifact_context_is_bound_to_the_invoking_user() -> None:
    actor_id = "a" * 64
    prefix = (
        f"users/{actor_id}/bots/research-reports/artifacts/"
        "12345678-1234-1234-1234-123456789012"
    )
    assert (
        artifacts.artifact_prefix_from_payload(
            {"artifacts": {"prefix": prefix}}, actor_id
        )
        == prefix
    )
    with pytest.raises(ValueError, match="invoking user"):
        artifacts.artifact_prefix_from_payload(
            {"artifacts": {"prefix": prefix}}, "b" * 64
        )


def test_group_artifact_context_requires_group_scope() -> None:
    prefix = (
        "groups/12345678-1234-1234-1234-123456789012/artifacts/"
        "22345678-1234-1234-1234-123456789012"
    )
    assert (
        artifacts.artifact_prefix_from_payload(
            {"group": {}, "artifacts": {"prefix": prefix}}
        )
        == prefix
    )
    with pytest.raises(ValueError, match="group artifact scope"):
        artifacts.artifact_prefix_from_payload({"artifacts": {"prefix": prefix}})


@pytest.mark.parametrize("scope", ["personal", "group"])
def test_group_artifacts_reject_a_different_groups_or_personal_actor(scope) -> None:
    group_id = "12345678-1234-1234-1234-123456789012"
    prefix = f"groups/{group_id}/artifacts/{group_id}"
    actor_id = hashlib.sha256(b"group:another-group").hexdigest()
    with pytest.raises(ValueError, match="group artifact scope"):
        artifacts.artifact_prefix_from_payload(
            {
                "group": {},
                "artifacts": {"prefix": prefix},
                "memory": {"scope": scope, "actorId": actor_id},
            },
            actor_id,
        )


def test_group_artifacts_reject_a_group_actor_in_a_personal_scope() -> None:
    group_id = "12345678-1234-1234-1234-123456789012"
    actor_id = hashlib.sha256(f"group:{group_id}".encode()).hexdigest()
    with pytest.raises(ValueError, match="group artifact scope"):
        artifacts.artifact_prefix_from_payload(
            {
                "group": {},
                "artifacts": {"prefix": f"groups/{group_id}/artifacts/{group_id}"},
                "memory": {"scope": "personal", "actorId": actor_id},
            },
            actor_id,
        )


def test_artifact_tool_rejects_binary_and_oversized_outputs(monkeypatch) -> None:
    prefix = f"users/{'b' * 64}/artifacts/12345678123412341234123456789012"
    monkeypatch.setattr(artifacts, "FILES_BUCKET_NAME", "files")
    save = artifacts.artifact_tool(
        prefix, client=SimpleNamespace(put_object=lambda **_: None)
    )

    with pytest.raises(ValueError, match="must end"):
        save("archive.zip", "not really a zip")
    with pytest.raises(ValueError, match="UTF-8 bytes"):
        save("large.txt", "x" * (artifacts.MAX_ARTIFACT_SOURCE_BYTES + 1))


@pytest.mark.parametrize(
    ("filename", "content", "content_type", "members"),
    [
        (
            "Launch notes.docx",
            "# Launch notes\n\nThe deployment is ready.\n\n- Runtime\n- Memory",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            {"word/document.xml"},
        ),
        (
            "Metrics.xlsx",
            "Metric,Value,As of\nRequests,42,2026-09-04\nSuccess rate,99.5%,2026-09-04",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            {"xl/workbook.xml", "xl/worksheets/sheet1.xml"},
        ),
        (
            "Plan.pptx",
            "# Release plan\n\n- Deploy\n- Verify\n---\n# Follow up\n\nReview service metrics.",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            {"ppt/presentation.xml", "ppt/slides/slide1.xml"},
        ),
    ],
)
def test_artifact_tool_renders_native_office_files(
    monkeypatch, filename, content, content_type, members
) -> None:
    prefix = f"users/{'c' * 64}/artifacts/12345678-1234-1234-1234-123456789012"
    target = FakeS3()
    monkeypatch.setattr(artifacts, "FILES_BUCKET_NAME", "files")
    save = artifacts.artifact_tool(prefix, client=target)

    save(filename, content)

    request = target.requests[0]
    assert request["ContentType"] == content_type
    with zipfile.ZipFile(io.BytesIO(request["Body"])) as archive:
        assert members <= set(archive.namelist())


def test_artifact_tool_renders_a_real_pdf(monkeypatch) -> None:
    prefix = f"users/{'d' * 64}/artifacts/12345678-1234-1234-1234-123456789012"
    target = FakeS3()
    monkeypatch.setattr(artifacts, "FILES_BUCKET_NAME", "files")
    save = artifacts.artifact_tool(prefix, client=target)

    save("Deployment report.pdf", "# Deployment report\n\nAll checks passed.")

    request = target.requests[0]
    assert request["ContentType"] == "application/pdf"
    assert request["Body"].startswith(b"%PDF-")
    assert request["Body"].endswith(b"%%EOF\n")


def test_spreadsheet_formula_like_values_remain_plain_text(monkeypatch) -> None:
    prefix = f"users/{'f' * 64}/artifacts/12345678-1234-1234-1234-123456789012"
    target = FakeS3()
    monkeypatch.setattr(artifacts, "FILES_BUCKET_NAME", "files")
    save = artifacts.artifact_tool(prefix, client=target)

    save("Safe export.xlsx", "Value\n=1+1\n+SUM(A1:A2)\n@danger")

    workbook = load_workbook(io.BytesIO(target.requests[0]["Body"]), data_only=False)
    assert [
        workbook.active.cell(row=index, column=1).data_type for index in range(2, 5)
    ] == [
        "s",
        "s",
        "s",
    ]
