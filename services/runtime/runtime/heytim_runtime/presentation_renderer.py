from __future__ import annotations

import io
import re

from .artifact_content import (
    ContentBlock,
    document_title,
    markdown_blocks,
    validate_content,
)

MAX_SLIDES = 40


def presentation_slides(
    content: str, filename: str
) -> list[tuple[str, list[ContentBlock]]]:
    sections = re.split(r"(?m)^\s*---\s*$", validate_content(content))
    slides: list[tuple[str, list[ContentBlock]]] = []
    for section in sections:
        if not section.strip():
            continue
        blocks = markdown_blocks(section)
        if blocks and blocks[0].kind == "heading":
            title = blocks[0].text
            body = blocks[1:]
        else:
            title = document_title(filename) if not slides else "Continued"
            body = blocks
        chunks: list[list[ContentBlock]] = [[]]
        characters = 0
        for block in body:
            if chunks[-1] and (
                len(chunks[-1]) >= 8 or characters + len(block.text) > 900
            ):
                chunks.append([])
                characters = 0
            chunks[-1].append(block)
            characters += len(block.text)
        for index, chunk in enumerate(chunks):
            chunk_title = title if index == 0 else f"{title} continued"
            slides.append((chunk_title, chunk))
    if len(slides) > MAX_SLIDES:
        raise ValueError(f"presentation content may create at most {MAX_SLIDES} slides")
    return slides


def render_pptx(filename: str, content: str) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    slides = presentation_slides(content, filename)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    for title, blocks in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(250, 250, 247)
        title_shape = slide.shapes.title
        title_shape.text = title
        title_paragraph = title_shape.text_frame.paragraphs[0]
        title_paragraph.font.name = "Arial"
        title_paragraph.font.size = Pt(32)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = RGBColor(0, 0, 0)

        body_shape = slide.placeholders[1]
        body_shape.left = Inches(0.9)
        body_shape.top = Inches(1.65)
        body_shape.width = Inches(11.5)
        body_shape.height = Inches(4.9)
        text_frame = body_shape.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        font_size = 18 if sum(len(block.text) for block in blocks) > 620 else 21
        for index, block in enumerate(blocks):
            paragraph = (
                text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            )
            paragraph.text = block.text
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(font_size)
            paragraph.font.color.rgb = RGBColor(35, 35, 32)
            paragraph.space_after = Pt(10)
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph_properties = paragraph._p.get_or_add_pPr()
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buBlip"):
                existing = paragraph_properties.find(qn(tag))
                if existing is not None:
                    paragraph_properties.remove(existing)
            if block.kind == "bullet":
                marker = OxmlElement("a:buChar")
                marker.set("char", "•")
                paragraph_properties.insert(0, marker)
                paragraph_properties.set("marL", "342900")
                paragraph_properties.set("indent", "-285750")
            elif block.kind == "number":
                marker = OxmlElement("a:buAutoNum")
                marker.set("type", "arabicPeriod")
                paragraph_properties.insert(0, marker)
                paragraph_properties.set("marL", "342900")
                paragraph_properties.set("indent", "-285750")
            else:
                paragraph_properties.insert(0, OxmlElement("a:buNone"))
                paragraph_properties.set("marL", "0")
                paragraph_properties.set("indent", "0")
            if block.kind == "heading":
                paragraph.font.bold = True
                paragraph.font.size = Pt(font_size + 2)
            elif block.kind == "code":
                paragraph.font.name = "Courier New"
                paragraph.font.size = Pt(max(17, font_size - 2))

    presentation.core_properties.title = slides[0][0]
    presentation.core_properties.author = "HeyTim"
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()
