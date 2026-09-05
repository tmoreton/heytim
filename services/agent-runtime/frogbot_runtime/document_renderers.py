from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from .artifact_content import document_title, markdown_blocks


def render_docx(filename: str, content: str) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_BREAK, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    blocks = markdown_blocks(content)
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    for style_name, size, bold in (
        ("Normal", 11, False),
        ("Title", 24, True),
        ("Heading 1", 17, True),
        ("Heading 2", 14, True),
        ("Heading 3", 12, True),
    ):
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        paragraph_properties = style.element.pPr
        if paragraph_properties is not None:
            border = paragraph_properties.find(qn("w:pBdr"))
            if border is not None:
                paragraph_properties.remove(border)
    normal = document.styles["Normal"]
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    if "FroggyBot Code" not in document.styles:
        code_style = document.styles.add_style(
            "FroggyBot Code", WD_STYLE_TYPE.PARAGRAPH
        )
        code_style.font.name = "Courier New"
        code_style.font.size = Pt(9)
        code_style.element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
        code_style.paragraph_format.space_after = Pt(8)

    first_heading = blocks[0] if blocks and blocks[0].kind == "heading" else None
    if first_heading:
        document.add_paragraph(first_heading.text, style="Title")
        blocks = blocks[1:]
    else:
        document.add_paragraph(document_title(filename), style="Title")

    for block in blocks:
        if block.kind == "heading":
            document.add_paragraph(
                block.text, style=f"Heading {min(3, max(1, block.level))}"
            )
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "number":
            document.add_paragraph(block.text, style="List Number")
        elif block.kind == "code":
            paragraph = document.add_paragraph(style="FroggyBot Code")
            for index, line in enumerate(block.text.split("\n")):
                if index:
                    paragraph.add_run().add_break(WD_BREAK.LINE)
                paragraph.add_run(line)
        else:
            document.add_paragraph(block.text)

    document.core_properties.title = (
        first_heading.text if first_heading else document_title(filename)
    )
    document.core_properties.author = "FroggyBot"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _reportlab_fonts() -> tuple[str, str, str]:
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    fonts_dir = Path(reportlab.__file__).resolve().parent / "fonts"
    regular = fonts_dir / "Vera.ttf"
    bold = fonts_dir / "VeraBd.ttf"
    mono = fonts_dir / "VeraMono.ttf"
    if regular.is_file() and bold.is_file() and mono.is_file():
        for name, path in (
            ("FroggyBotSans", regular),
            ("FroggyBotSansBold", bold),
            ("FroggyBotMono", mono),
        ):
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(path)))
        return "FroggyBotSans", "FroggyBotSansBold", "FroggyBotMono"
    return "Helvetica", "Helvetica-Bold", "Courier"


def render_pdf(filename: str, content: str) -> bytes:
    from html import escape

    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, Preformatted, SimpleDocTemplate, Spacer

    regular, bold, mono = _reportlab_fonts()
    blocks = markdown_blocks(content)
    first_heading = blocks[0] if blocks and blocks[0].kind == "heading" else None
    title = first_heading.text if first_heading else document_title(filename)
    if first_heading:
        blocks = blocks[1:]

    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        title=title,
        author="FroggyBot",
        leftMargin=0.8 * inch,
        rightMargin=0.8 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.7 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "FroggyBotTitle",
        parent=styles["Title"],
        fontName=bold,
        fontSize=22,
        leading=27,
        textColor=HexColor("#000000"),
        alignment=TA_CENTER,
        spaceAfter=22,
    )
    body_style = ParagraphStyle(
        "FroggyBotBody",
        parent=styles["BodyText"],
        fontName=regular,
        fontSize=10.5,
        leading=15,
        textColor=HexColor("#20231F"),
        spaceAfter=9,
    )
    bullet_style = ParagraphStyle(
        "FroggyBotBullet",
        parent=body_style,
        leftIndent=18,
        firstLineIndent=-9,
        bulletIndent=7,
    )
    heading_styles = {
        level: ParagraphStyle(
            f"FroggyBotHeading{level}",
            parent=styles[f"Heading{level}"],
            fontName=bold,
            fontSize=size,
            leading=size + 4,
            textColor=HexColor("#000000"),
            spaceBefore=12,
            spaceAfter=7,
        )
        for level, size in ((1, 17), (2, 14), (3, 12))
    }
    code_style = ParagraphStyle(
        "FroggyBotCode",
        parent=body_style,
        fontName=mono,
        fontSize=8.5,
        leading=12,
        leftIndent=12,
    )

    story: list[Any] = [Paragraph(escape(title, quote=False), title_style)]
    for block in blocks:
        text = escape(block.text, quote=False).replace("\n", "<br/>")
        if block.kind == "heading":
            story.append(Paragraph(text, heading_styles[min(3, max(1, block.level))]))
        elif block.kind == "bullet":
            story.append(Paragraph(text, bullet_style, bulletText="-"))
        elif block.kind == "number":
            story.append(Paragraph(text, bullet_style, bulletText="1."))
        elif block.kind == "code":
            story.append(Preformatted(block.text, code_style))
        else:
            story.append(Paragraph(text, body_style))
    story.append(Spacer(1, 1))

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(regular, 8)
        canvas.setFillColor(HexColor("#777777"))
        canvas.drawCentredString(letter[0] / 2, 0.35 * inch, str(doc.page))
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()
